#!/usr/bin/env python3
"""Build the editable advisor-review DOCX from the consolidated Markdown draft.

This is a document-production helper only. It does not read or alter experiment
artifacts. The generated DOCX embeds the five methodological figures and leaves
the three qualitative result figures explicitly reserved until provenance-backed
images become available.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


AUTHOR = "João Francisco de Souza Santos Lemos"
TITLE = "Reconstrução neural de cenas domésticas estruturadas para simulação robótica"
QUALIFICATION_TITLE = (
    "Reconstrução Semântica Automatizada de Ambientes Domésticos para Simulação de Robôs"
)
ACCENT = "1F4E79"
MUTED = "5B6573"
TOC_ENTRIES = (
    ("1 Introdução", 3),
    ("2 Trabalhos relacionados", 7),
    ("3 Metodologia", 12),
    ("4 Protocolo experimental", 19),
    ("5 Resultados", 23),
    ("6 Discussão preliminar", 27),
    ("7 Limitações metodológicas", 29),
    ("8 Considerações finais", 30),
    ("Referências", 31),
)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep])
    if placeholder:
        run._r.append(OxmlElement("w:t"))
        run._r[-1].text = placeholder
    run._r.append(fld_end)


def add_page_number(section) -> None:
    footer = section.footer
    footer.is_linked_to_previous = True
    paragraph = footer.paragraphs[0]
    if footer._element.xpath(".//w:fldChar"):
        return
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.style = "Footer"
    add_field(paragraph, " PAGE ", "1")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


INLINE_TOKEN = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*]+?\*)")


def add_inline(paragraph, text: str) -> None:
    cursor = 0
    for match in INLINE_TOKEN.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Liberation Mono"
            run.font.size = Pt(9.5)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Liberation Serif"
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Cm(1.25)

    for name, size, before, after in (
        ("Title", 20, 0, 18),
        ("Heading 1", 16, 18, 10),
        ("Heading 2", 14, 14, 8),
        ("Heading 3", 12, 12, 6),
    ):
        style = styles[name]
        style.font.name = "Liberation Sans"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(ACCENT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Cm(0)

    caption = styles["Caption"]
    caption.font.name = "Liberation Serif"
    caption.font.size = Pt(10)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.first_line_indent = Cm(0)

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Liberation Serif"
        style.font.size = Pt(12)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        style.paragraph_format.space_after = Pt(3)

    if "Editorial Note" not in styles:
        note = styles.add_style("Editorial Note", WD_STYLE_TYPE.PARAGRAPH)
        note.font.name = "Liberation Serif"
        note.font.size = Pt(10.5)
        note.font.italic = True
        note.font.color.rgb = RGBColor.from_string(MUTED)
        note.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        note.paragraph_format.space_after = Pt(10)
        note.paragraph_format.first_line_indent = Cm(0)


def configure_section(section, landscape: bool = False) -> None:
    section.different_first_page_header_footer = False
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.header_distance = Cm(1)
    section.footer_distance = Cm(1)
    add_page_number(section)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(45)
    r = p.add_run(AUTHOR.upper())
    r.bold = True
    r.font.name = "Liberation Sans"
    r.font.size = Pt(13)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(95)
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run(TITLE)
    r.bold = True
    r.font.name = "Liberation Sans"
    r.font.size = Pt(21)
    r.font.color.rgb = RGBColor.from_string(ACCENT)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("Título provisório")
    r.italic = True
    r.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(45)
    p.paragraph_format.left_indent = Cm(4)
    p.paragraph_format.right_indent = Cm(4)
    r = p.add_run("Dissertação de Mestrado — versão preliminar para revisão do orientador")
    r.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    p.add_run("Outubro de 2026")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    r = p.add_run(f'Título utilizado na qualificação: “{QUALIFICATION_TITLE}”.')
    r.italic = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor.from_string(MUTED)


def add_toc(doc: Document) -> None:
    doc.add_page_break()
    heading = doc.add_paragraph("Sumário", style="Title")
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    table = doc.add_table(rows=len(TOC_ENTRIES), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row, (label, page_number) in zip(table.rows, TOC_ENTRIES):
        row.cells[0].width = Cm(14)
        row.cells[1].width = Cm(1)
        for cell in row.cells:
            set_cell_margins(cell, top=55, start=0, bottom=55, end=0)
        left = row.cells[0].paragraphs[0]
        left.paragraph_format.first_line_indent = Cm(0)
        left.paragraph_format.space_after = Pt(0)
        add_inline(left, label)
        right = row.cells[1].paragraphs[0]
        right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        right.paragraph_format.first_line_indent = Cm(0)
        right.paragraph_format.space_after = Pt(0)
        right.add_run(str(page_number))
    doc.add_page_break()


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        for col_index in range(columns):
            cell = row.cells[col_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            text = values[col_index] if col_index < len(values) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if columns <= 3 or col_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            )
            add_inline(paragraph, text)
            for run in paragraph.runs:
                run.font.name = "Liberation Sans"
                run.font.size = Pt(7.5 if columns >= 7 else 9)
            if row_index == 0:
                set_cell_shading(cell, "D9EAF4")
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(ACCENT)
        if row_index == 0:
            set_repeat_table_header(row)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    raw_rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        parts = [part.strip() for part in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", part) for part in parts):
            raw_rows.append(parts)
        index += 1
    return raw_rows, index


def add_source_content(doc: Document, source: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    index = 0
    first_chapter = True
    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            index += 1
            continue

        if stripped.startswith("# "):
            index += 1
            continue

        if stripped.startswith("## "):
            paragraph = doc.add_paragraph(stripped[3:].strip(), style="Heading 1")
            if not first_chapter:
                paragraph.paragraph_format.page_break_before = True
            first_chapter = False
            index += 1
            continue

        if stripped.startswith("### "):
            doc.add_paragraph(stripped[4:].strip(), style="Heading 2")
            index += 1
            continue

        if stripped.startswith("#### "):
            doc.add_paragraph(stripped[5:].strip(), style="Heading 3")
            index += 1
            continue

        image_match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", stripped)
        if image_match:
            image_path = (source.parent / image_match.group(2)).resolve()
            if image_path.suffix.lower() == ".svg":
                png_candidate = image_path.with_suffix(".png")
                if png_candidate.exists():
                    image_path = png_candidate
            if not image_path.exists():
                raise FileNotFoundError(f"Figure not found: {image_path}")
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.keep_with_next = True
            run = paragraph.add_run()
            width = 5.6 if image_path.stem == "figure_03_floor" else 6.25
            run.add_picture(str(image_path), width=Inches(width))
            index += 1
            continue

        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            comparison_table = bool(rows and rows[0] and rows[0][0] == "Método")
            if comparison_table:
                landscape = doc.add_section(WD_SECTION.NEW_PAGE)
                configure_section(landscape, landscape=True)
            add_markdown_table(doc, rows)
            if comparison_table:
                while index < len(lines) and not lines[index].strip():
                    index += 1
                if index < len(lines):
                    caption = lines[index].strip()
                    if re.fullmatch(r"\*Tabela \d+.+\*", caption):
                        paragraph = doc.add_paragraph(style="Caption")
                        add_inline(paragraph, caption[1:-1])
                        index += 1
                portrait = doc.add_section(WD_SECTION.NEW_PAGE)
                configure_section(portrait, landscape=False)
            continue

        if re.fullmatch(r"\*Figura \d+.+\*", stripped) or re.fullmatch(
            r"\*Tabela \d+.+\*", stripped
        ):
            paragraph = doc.add_paragraph(style="Caption")
            add_inline(paragraph, stripped[1:-1])
            index += 1
            continue

        if stripped.startswith(">"):
            text = stripped.lstrip("> ").replace("**", "")
            paragraph = doc.add_paragraph(style="Editorial Note")
            paragraph.paragraph_format.left_indent = Cm(1)
            paragraph.paragraph_format.right_indent = Cm(1)
            add_inline(paragraph, text)
            index += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            paragraph = doc.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+\.\s+", "", stripped))
            index += 1
            continue

        if re.match(r"^[-*]\s+", stripped):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, re.sub(r"^[-*]\s+", "", stripped))
            index += 1
            continue

        if stripped.startswith("*") and stripped.endswith("*"):
            paragraph = doc.add_paragraph(style="Editorial Note")
            add_inline(paragraph, stripped[1:-1])
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate:
                index += 1
                break
            if candidate.startswith(("#", ">", "|", "![")):
                break
            if re.match(r"^(\d+\.|[-*])\s+", candidate):
                break
            if (candidate.startswith("*") and candidate.endswith("*")):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = doc.add_paragraph()
        add_inline(paragraph, " ".join(paragraph_lines))


def build(source: Path, output: Path) -> None:
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])
    doc.sections[0].different_first_page_header_footer = True
    doc.core_properties.title = TITLE
    doc.core_properties.author = AUTHOR
    doc.core_properties.subject = "Versão preliminar para revisão do orientador"
    doc.core_properties.keywords = "DM-NeRF; DM-SR; reconstrução neural; Gazebo Classic"

    settings = doc.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)

    add_cover(doc)
    add_toc(doc)
    add_source_content(doc, source)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.source.resolve(), args.output.resolve())


if __name__ == "__main__":
    main()
